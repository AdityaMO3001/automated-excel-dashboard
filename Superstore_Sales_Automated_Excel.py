import pandas as pd
import openpyxl
from openpyxl.chart import LineChart, Reference
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders

# Step 1: Load and Clean the Data
file_path = 'superstore_sales.xlsx'  # Ensure you have the correct file path
df = pd.read_excel(file_path)

# Display basic info about the dataset
print(df.info())

# Clean the data by dropping rows with missing sales or profit
df = df.dropna(subset=['Sales', 'Profit'])

# Ensure 'Order Date' is in datetime format
df['Order Date'] = pd.to_datetime(df['Order Date'])

# Create new columns for 'Year' and 'Month' for easy grouping
df['Year'] = df['Order Date'].dt.year
df['Month'] = df['Order Date'].dt.month_name()

# Step 2: Aggregate the data by Year and Month
monthly_sales = df.groupby(['Year', 'Month'])[['Sales', 'Profit']].sum().reset_index()

# Step 3: Create an Excel Dashboard with Charts
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Sales Dashboard'

# Add headers for the data
headers = ['Year', 'Month', 'Total Sales', 'Total Profit']
ws.append(headers)

# Add the aggregated data to the Excel sheet
for index, row in monthly_sales.iterrows():
    ws.append(row.tolist())

# Create a Line Chart for Sales over time
chart_sales = LineChart()
data_sales = Reference(ws, min_col=3, min_row=1, max_col=3, max_row=len(monthly_sales)+1)
chart_sales.add_data(data_sales, titles_from_data=True)
chart_sales.title = "Sales Over Time"
ws.add_chart(chart_sales, "F5")

# Create a Line Chart for Profit over time
chart_profit = LineChart()
data_profit = Reference(ws, min_col=4, min_row=1, max_col=4, max_row=len(monthly_sales)+1)
chart_profit.add_data(data_profit, titles_from_data=True)
chart_profit.title = "Profit Over Time"
ws.add_chart(chart_profit, "F20")

# Save the Excel dashboard
dashboard_file = 'sales_dashboard.xlsx'
wb.save(dashboard_file)

# Step 4: Automate Sending the Dashboard via Email
def send_email_with_attachment():
    from_addr = 'mogadpally.a@northeastern.edu'  # Replace with your email
    to_addr = 'mogadpallya@gmail.com'  # Replace with the recipient's email
    subject = "Automated Sales Dashboard"
    body = "Please find the latest sales dashboard attached."

    # Create the email
    msg = MIMEMultipart()
    msg['From'] = from_addr
    msg['To'] = to_addr
    msg['Subject'] = subject
    msg.attach(MIMEText(body, 'plain'))

    # Attach the Excel file
    attachment = open(dashboard_file, 'rb')
    part = MIMEBase('application', 'octet-stream')
    part.set_payload(attachment.read())
    encoders.encode_base64(part)
    part.add_header('Content-Disposition', f'attachment; filename={dashboard_file}')
    msg.attach(part)

    # Send the email
    try:
        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(from_addr, 'Aditya@123')  # Replace with your email password
        text = msg.as_string()
        server.sendmail(from_addr, to_addr, text)
        server.quit()
        print("Email sent successfully!")
    except Exception as e:
        print(f"Error: {str(e)}")

# Call the function to send the email
send_email_with_attachment()
